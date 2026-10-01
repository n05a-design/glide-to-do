from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


path = Path(sys.argv[1])
doc = Document(path)

sections = []
for i, section in enumerate(doc.sections, 1):
    sections.append(
        {
            "section": i,
            "page_width_in": round(section.page_width.inches, 3),
            "page_height_in": round(section.page_height.inches, 3),
            "top_margin_in": round(section.top_margin.inches, 3),
            "bottom_margin_in": round(section.bottom_margin.inches, 3),
            "left_margin_in": round(section.left_margin.inches, 3),
            "right_margin_in": round(section.right_margin.inches, 3),
        }
    )

tables = []
for i, table in enumerate(doc.tables, 1):
    first = table.rows[0]
    first_trpr = first._tr.get_or_add_trPr()
    repeat_header = first_trpr.find(qn("w:tblHeader")) is not None
    rows = []
    for r_idx, row in enumerate(table.rows, 1):
        trpr = row._tr.get_or_add_trPr()
        rows.append(
            {
                "row": r_idx,
                "cant_split": trpr.find(qn("w:cantSplit")) is not None,
                "text": " | ".join(" ".join(cell.text.split()) for cell in row.cells)[:260],
            }
        )
    tables.append(
        {
            "table": i,
            "rows": len(table.rows),
            "columns": len(table.columns),
            "autofit": table.autofit,
            "repeat_first_row": repeat_header,
            "header": rows[0]["text"],
            "rows_with_cant_split": [r["row"] for r in rows if r["cant_split"]],
            "last_row": rows[-1]["text"],
        }
    )

headings = []
for idx, p in enumerate(doc.paragraphs, 1):
    if p.style and p.style.name.lower().startswith("heading"):
        headings.append({"paragraph": idx, "style": p.style.name, "text": " ".join(p.text.split())})

print(json.dumps({"sections": sections, "tables": tables, "headings": headings}, ensure_ascii=False, indent=2))
