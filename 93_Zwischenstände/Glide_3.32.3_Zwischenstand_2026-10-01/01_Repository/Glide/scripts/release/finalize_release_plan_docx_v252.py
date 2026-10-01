from __future__ import annotations

import argparse
import os
import re
import tempfile
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from pypdf import PdfReader


TOC_BOOKMARK_PREFIX = "GlideToc"
EXTENDED_PROPERTIES_NS = "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"
TOC_DISPLAY_LABELS = {
    4: "4. Ergänzungen und offene Vorbereitungspunkte",
    7: "7. Aufgabenverteilung Codex / Benutzer",
    12: "12. Apple-/Microsoft-Accountentscheidungen",
    13: "13. Reproduzierbares QA-Testsystem",
    18: "18. Nachgeholte Arbeitsvorbereitung",
    21: "21. Zielzustand Release-Workflow",
    22: "22. Fazit und nächste Schritte",
}
HISTORICAL_SCREENSHOT_ALT_TEXT = (
    "Historischer Windows-Explorer-Screenshot des Glide-Arbeitsstands vom 31.08.2026: "
    "Ordner für Probelisten, Python-Versionen, Dokumentation und Archiv sowie die "
    "damalige Testdatei _test_glide_v250.py."
)


def normalize_text(value: str) -> str:
    value = value.replace("\u00ad", "").replace("–", "-").replace("—", "-")
    return re.sub(r"\s+", " ", value).strip().casefold()


def heading_paragraphs(document):
    headings = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.style is not None
        and paragraph.style.name == "Heading 1"
        and re.match(r"^\d+\.\s", paragraph.text)
    ]
    if len(headings) != 23:
        raise ValueError(f"Erwartet wurden 23 Hauptkapitel, gefunden: {len(headings)}")
    return headings


def extract_heading_pages(document, pdf_path: Path) -> tuple[dict[str, int], int]:
    reader = PdfReader(str(pdf_path))
    page_texts = [normalize_text(page.extract_text() or "") for page in reader.pages]
    mapping: dict[str, int] = {}
    for heading in heading_paragraphs(document):
        needle = normalize_text(heading.text)
        # Seite 2 enthält die Inhaltsübersicht selbst und damit ebenfalls alle
        # Kapiteltexte. Für die Zielseite zählt erst der eigentliche Kapitelkopf
        # ab Seite 3.
        matches = [index + 1 for index, page_text in enumerate(page_texts) if index >= 2 and needle in page_text]
        if len(matches) != 1:
            raise ValueError(f"Kapitel konnte nicht eindeutig im PDF gefunden werden: {heading.text!r} -> {matches}")
        mapping[heading.text] = matches[0]
    return mapping, len(reader.pages)


def clear_paragraph(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def add_hyperlink_run(hyperlink, text: str, *, tab: bool = False) -> None:
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{attribute}"), "Arial")
    r_pr.append(r_fonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "17")
    r_pr.append(size)
    size_cs = OxmlElement("w:szCs")
    size_cs.set(qn("w:val"), "17")
    r_pr.append(size_cs)
    run.append(r_pr)
    if tab:
        run.append(OxmlElement("w:tab"))
    else:
        node = OxmlElement("w:t")
        node.text = text
        run.append(node)
    hyperlink.append(run)


def add_internal_link(paragraph, text: str, bookmark: str, page: int) -> None:
    clear_paragraph(paragraph)
    paragraph.paragraph_format.space_after = Pt(1.5)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        Cm(7.55),
        WD_TAB_ALIGNMENT.RIGHT,
        WD_TAB_LEADER.DOTS,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), bookmark)
    hyperlink.set(qn("w:history"), "1")
    add_hyperlink_run(hyperlink, text)
    add_hyperlink_run(hyperlink, "", tab=True)
    add_hyperlink_run(hyperlink, str(page))
    paragraph._p.append(hyperlink)


def toc_display_label(index: int, heading_text: str) -> str:
    """Keep the TOC compact without changing the actual heading/bookmark text."""
    return TOC_DISPLAY_LABELS.get(index, heading_text)


def rebuild_toc(document, pages: dict[str, int]) -> None:
    headings = heading_paragraphs(document)
    entries = [
        (
            toc_display_label(index, heading.text),
            f"{TOC_BOOKMARK_PREFIX}{index:02d}",
            pages[heading.text],
        )
        for index, heading in enumerate(headings, start=1)
    ]
    toc_table = document.tables[1]
    columns = (entries[:12], entries[12:])
    for cell, cell_entries in zip(toc_table.rows[0].cells, columns):
        first = cell.paragraphs[0]
        for extra in list(cell.paragraphs[1:]):
            extra._element.getparent().remove(extra._element)
        for entry_index, (label, bookmark, page) in enumerate(cell_entries):
            paragraph = first if entry_index == 0 else cell.add_paragraph()
            add_internal_link(paragraph, label, bookmark, page)


def style_footer_run(run) -> None:
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)


def add_page_field(paragraph) -> None:
    begin_run = paragraph.add_run()
    style_footer_run(begin_run)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin_run._r.append(begin)

    instruction_run = paragraph.add_run()
    style_footer_run(instruction_run)
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    instruction_run._r.append(instruction)

    separate_run = paragraph.add_run()
    style_footer_run(separate_run)
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    separate_run._r.append(separate)

    result_run = paragraph.add_run("1")
    style_footer_run(result_run)

    end_run = paragraph.add_run()
    style_footer_run(end_run)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run._r.append(end)


def rebuild_footer(document) -> None:
    """Use one tab-aligned line so narrow renderers cannot clip table cells."""
    visited: set[int] = set()
    for section in document.sections:
        # Word can keep independent stories for default, even-page and
        # first-page footers. Rebuild all three so odd/even pagination renders
        # identically in Word and LibreOffice.
        for footer in (section.footer, section.even_page_footer, section.first_page_footer):
            footer_key = id(footer._element)
            if footer_key in visited:
                continue
            visited.add(footer_key)

            for table in list(footer.tables):
                footer._element.remove(table._element)

            paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
            for extra in list(footer.paragraphs[1:]):
                footer._element.remove(extra._element)
            clear_paragraph(paragraph)
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            writable_width = section.page_width - section.left_margin - section.right_margin
            paragraph.paragraph_format.tab_stops.add_tab_stop(
                writable_width,
                WD_TAB_ALIGNMENT.RIGHT,
            )

            left = paragraph.add_run("Technische Arbeitsgrundlage · Stand 01.09.2026")
            style_footer_run(left)
            tab = paragraph.add_run("\t")
            style_footer_run(tab)
            page_prefix = paragraph.add_run("Seite ")
            style_footer_run(page_prefix)
            add_page_field(paragraph)


def set_image_alt_text(document) -> None:
    image_properties = list(document._element.iter(qn("wp:docPr")))
    if len(image_properties) != 1:
        raise ValueError(f"Erwartet wurde genau eine eingebettete Abbildung, gefunden: {len(image_properties)}")
    image_properties[0].set("title", "Historischer Glide-Arbeitsstand")
    image_properties[0].set("descr", HISTORICAL_SCREENSHOT_ALT_TEXT)


def document_statistics(document) -> dict[str, int]:
    texts: list[str] = []
    paragraph_count = 0
    line_count = 0

    def add_paragraph(paragraph) -> None:
        nonlocal paragraph_count, line_count
        text = paragraph.text.strip()
        if not text:
            return
        texts.append(text)
        paragraph_count += 1
        line_count += text.count("\n") + 1

    for paragraph in document.paragraphs:
        add_paragraph(paragraph)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    add_paragraph(paragraph)

    combined = "\n".join(texts)
    words = re.findall(r"\b[\wÄÖÜäöüß]+(?:[-'][\wÄÖÜäöüß]+)*\b", combined, flags=re.UNICODE)
    return {
        "Words": len(words),
        "Characters": len(re.sub(r"\s", "", combined)),
        "CharactersWithSpaces": len(combined),
        "Paragraphs": paragraph_count,
        "Lines": line_count,
    }


def update_extended_properties(path: Path, pages: int, stats: dict[str, int]) -> None:
    ET.register_namespace("", EXTENDED_PROPERTIES_NS)
    with zipfile.ZipFile(path, "r") as source_zip:
        entries = [(info, source_zip.read(info.filename)) for info in source_zip.infolist()]

    updated_entries = []
    for info, data in entries:
        if info.filename == "docProps/app.xml":
            root = ET.fromstring(data)
            values = {"Pages": pages, **stats}
            for name, value in values.items():
                element = root.find(f"{{{EXTENDED_PROPERTIES_NS}}}{name}")
                if element is None:
                    element = ET.SubElement(root, f"{{{EXTENDED_PROPERTIES_NS}}}{name}")
                element.text = str(value)
            data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
        updated_entries.append((info, data))

    with tempfile.NamedTemporaryFile(prefix="glide-docx-stats-", suffix=".docx", dir=path.parent, delete=False) as handle:
        temporary = Path(handle.name)
    try:
        with zipfile.ZipFile(temporary, "w") as target_zip:
            for info, data in updated_entries:
                target_zip.writestr(info, data)
        os.replace(temporary, path)
    finally:
        if temporary.exists():
            temporary.unlink()


def finalize(source: Path, rendered_pdf: Path, output: Path) -> None:
    document = Document(source)
    pages, page_count = extract_heading_pages(document, rendered_pdf)
    rebuild_toc(document, pages)
    rebuild_footer(document)
    set_image_alt_text(document)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    update_extended_properties(output, page_count, document_statistics(document))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("rendered_pdf", type=Path)
    parser.add_argument("output", type=Path)
    arguments = parser.parse_args()
    finalize(arguments.source, arguments.rendered_pdf, arguments.output)


if __name__ == "__main__":
    main()
