"""Bearbeitbarer, quellengebundener Prüfbericht als DOCX; gebündeltes Python verwenden."""
from pathlib import Path
import hashlib
import json
from datetime import datetime, timezone
from collections import Counter
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).resolve().parent
DOCX = OUT / 'Glide_Abnahmebericht_3.49.0_2026-10-28.docx'
data = json.loads((OUT/'paketmatrix.json').read_text(encoding='utf-8'))
tests = json.loads((OUT/'testergebnis.json').read_text(encoding='utf-8'))
freeze = json.loads((OUT/'eingefroren.json').read_text(encoding='utf-8'))
code = json.loads((OUT/'codeinventar.json').read_text(encoding='utf-8'))
delivery = json.loads((OUT/'lieferstand.json').read_text(encoding='utf-8'))
counts = json.loads((OUT/'testzaehlung.json').read_text(encoding='utf-8'))
plan = json.loads((OUT/'planregister.json').read_text(encoding='utf-8'))
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(1.75), Cm(1.75)
sec.left_margin, sec.right_margin = Cm(2), Cm(2)
sec.header_distance, sec.footer_distance = Cm(.7), Cm(.7)
for style in ('Normal','Body Text'):
    doc.styles[style].font.name='Calibri'
    doc.styles[style].font.size=Pt(10.5)
    doc.styles[style].paragraph_format.space_after=Pt(6)
    doc.styles[style].paragraph_format.line_spacing=1.08
for name,size in [('Title',24),('Subtitle',11),('Heading 1',16),('Heading 2',12)]:
    doc.styles[name].font.name='Calibri'
    doc.styles[name].font.size=Pt(size)
    doc.styles[name].font.color.rgb=RGBColor(0,0,0)
    doc.styles[name].font.underline=False
    doc.styles[name].paragraph_format.space_after=Pt(8)
doc.core_properties.title='Glide Prüfbericht zur Abnahme der begonnenen Arbeiten'
doc.core_properties.subject='Stand 3.49.0 vom 28.10.2026 mit getrennten Prüf und Liefergrenzen'
doc.core_properties.author='Glide Projektprüfung'
doc.core_properties.keywords='Abnahme, Prüfkandidat, 3.49.0, WA-41, offene Arbeiten'
doc.core_properties.created=doc.core_properties.modified=datetime.now(timezone.utc)
doc.core_properties.last_modified_by='Glide Projektprüfung'
# Der Standardvorlage entstammende blaue Titelregeln entfernen.
for border in list(doc.styles.element.xpath('.//w:pBdr')):
    border.getparent().remove(border)
header=sec.header.paragraphs[0]
header.text='GLIDE  |  Prüfbericht zur Arbeitsabnahme  |  28.10.2026'
header.style=doc.styles['Normal']
header.runs[0].font.size=Pt(8)
header.runs[0].font.color.rgb=RGBColor(0,0,0)
footer=sec.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Stand 3.49.0  ·  Seite ').font.size=Pt(8)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)


def clean(text):
    text=str(text)
    for before,after in [('Kapitel6','Kapitel 6'),('Kapitel9','Kapitel 9'),('Format26','Format 26'),
                         ('Basis3.45','Basis 3.45'),('Quelle3.49','Quelle 3.49'),('auf3.','auf 3.'),
                         ('Kopf3.','Kopf 3.'),('Quellkandidaten3.','Quellkandidaten 3.'),
                         ('umfassenderen3.','umfassenderen 3.'),('aktuellen3.','aktuellen 3.'),('aktuelle3.','aktuelle 3.'),
                         ('Kernablauf für4.0','Kernablauf für 4.0'),('keine4.0','keine 4.0'),
                         ('Punkte1–4','Punkte 1–4'),('Stufe2','Stufe 2')]:
        text=text.replace(before,after)
    return text


def para(text='',bold=False):
    p=doc.add_paragraph()
    p.add_run(clean(text)).bold=bold
    return p


def heading(text,level=1):
    return doc.add_heading(text,level=level)


def page(title):
    doc.add_page_break()
    heading(title)


def table(headers,rows,widths=None,font=9.3):
    t=doc.add_table(rows=1,cols=len(headers))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    if widths is None: widths=[17/len(headers)]*len(headers)
    for i,w in enumerate(widths):t.columns[i].width=Cm(w)
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ('top','left','bottom','right','insideH','insideV'):
        el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
    pr.append(borders)
    for i,value in enumerate(headers):t.rows[0].cells[i].text=value
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for values in rows:
        cells=t.add_row().cells
        for i,value in enumerate(values):cells[i].text=clean(value)
    for ridx,row in enumerate(t.rows):
        no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
        for i,cell in enumerate(row.cells):
            cell.width=Cm(widths[i]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcp=cell._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for side in ('top','left','bottom','right'):
                edge=OxmlElement('w:'+side);edge.set(qn('w:w'),'75');edge.set(qn('w:type'),'dxa');margins.append(edge)
            tcp.append(margins)
            fill=OxmlElement('w:shd');fill.set(qn('w:val'),'clear');fill.set(qn('w:color'),'auto')
            fill.set(qn('w:fill'),'DCE6EF' if ridx==0 else ('F2F5F7' if ridx%2==0 else 'FFFFFF'));tcp.append(fill)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2);p.paragraph_format.line_spacing=1.02
                for run in p.runs:
                    run.font.name='Calibri';run.font.size=Pt(font)
                    if ridx==0:run.bold=True;run.font.color.rgb=RGBColor(0,0,0)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)
    return t


def matrix(title,rows,note):
    page(title)
    para(note)
    mapped=[]
    for entry in rows:
        state={'I':'Umgesetzt','T':'Teilweise','P':'Pausiert'}[entry['implementation']]
        mapped.append([entry['id']+'\n'+entry['title'],state+'\n'+entry['test'],entry['delivery'],entry['rest']])
    table(['Kennung und Umfang','Implementierung und Prüfung','Lieferstand','Offener Abschluss'],mapped,[4.6,3.9,2.1,6.4],9.1)


doc.add_paragraph('Glide Prüfbericht zur Abnahme',style='Title')
doc.add_paragraph('Begonnene Features und technische Arbeiten bis zum Quellkandidaten 3.49.0',style='Subtitle')
para('Prüfstichtag 28.10.2026  ·  Aufgabenformat 26  ·  Berichtsfassung 1')
heading('1 Ergebnis und Abnahmeempfehlung')
para('Nicht alle begonnenen Arbeiten sind abgeschlossen. Die gewählten Kernfunktionen sind überwiegend implementiert; eine vollständige Abnahme des Gesamtstands ist nicht belegt. Insbesondere P8/4.0, technische Querschnittsarbeiten, WA-41 und die Auslieferung des aktuellen Quellstands bleiben offen.',True)
para('Für einzelne belegte Implementierungen kann eine technische Teilabnahme mit ausdrücklich benannten Grenzen dokumentiert werden. Die Gesamtabnahme einschließlich aktueller Lieferpakete und öffentlicher Freigabe wird zurückgestellt. Dieser Bericht erteilt keine menschliche oder rechtsverbindliche Abnahmeentscheidung.')
table(['Prüfobjekt','Tatsächlicher Stand'],[
    ['Quellkandidat','3.49.0 / Format 26; Apphash im Quellenanhang'],
    ['Startbare Python-Fassung','3.47.0 / Format 26'],
    ['Eigenständige Windows-/Mac-Pakete','3.45.0; beide Archive aktuell hashgleich zum Liefernachweis'],
    ['Neue CI-Grundstufe','Bestanden; 19 Schritte, Lieferstand als Hinweis'],
    ['Neue Funktionsregression','FP-01: 8 von 9 Fällen grün, ein roter Gültigkeitsfall'],
    ['Breite Folgematrix','Nach rotem Fall gestoppt; zwölf Folgesuiten nicht ausgeführt'],
    ['Globale Abnahme','Offen: keine vollständig grüne aktuelle Voll-/Native-/Menschenabnahme']],[5.2,11.8],10)
heading('Unmittelbar abnahmerelevant',2)
para('1. Den FP-01-Prüfadapter an den sauberen Sitzungsabschluss anpassen und den unveränderten Pflichtumfang erneut bestehen.\n2. Begonnene Teilumfänge und dokumentierte Abnahmegrenzen gezielt schließen.\n3. WA-41A/B getrennt klären und anschließend aktuelle Vollmatrix, Lieferung und native Abnahmen belegen.')

page('2 Prüfmaßstab und tatsächlich ausgeführte Prüfung')
para('Der Abgleich umfasst den aktuellen Entwicklungsplan mit seinen historischen Aufträgen, das konsolidierte Funktionsinventar, Entscheidungen, Funktionsvertrag, QA-Nachweise, manuelle Prüfliste, Quellmodule und vorhandene Lieferartefakte. Die 52 Abnahmeeinheiten in den folgenden Matrizen überschneiden sich teilweise fachlich; sie sind keine addierbare Anzahl unabhängiger Produktfeatures. B01 bis B22 sind zusätzlich einzeln aufgeführt.')
para('Abschluss wird in vier getrennten Achsen bewertet: Implementierung, ausgeführte Prüfung, tatsächliche Lieferung sowie menschliche/native Abnahme. Eine vorhandene Funktion oder eine grüne Unit-Suite allein schließt die anderen Achsen nicht. Nicht gewählte Optionen gelten nicht als liegengebliebene Umsetzung.')
heading('Neue Nachweise dieses Audits',2)
table(['Prüfung','Ergebnis und Grenze'],[
    ['Unveränderliche Prüfkopie','Lokaler Gitklon mit aktuellen Änderungen; Quell-/Fixturehashes und isolierte Daten pro Prozess. Originalquelle nach dem Lauf unverändert.'],
    ['CI-Grundstufe','220 Werkzeugtests (ein Skip), zweimal 1.060 Fachtests, fünf Analysen, Syntax/Version/Fixtures, Startprobe sowie Datenschutz-/Ablage-/Synchronisationskontrollen bestanden.'],
    ['FP-01 Wiederholen','Acht Fälle grün; Gültigkeit rot mit TclError „Item … not found“. Kein grüner Gesamtstatus dieser Suite.'],
    ['Begrenzte Ursachenprobe','Originalfehler reproduziert. Diagnosevariation mit sauberem Test-Sitzungsabschluss besteht dieselben Assertions; keine Produktions- oder Testquelldatei geändert.'],
    ['Lieferkontrolle','Python 3.47.0/Format26 gelesen; Windows-/Mac-Archive 3.45.0 anhand SHA-256 und eingebetteter Versionsmetadaten geprüft. Keine neue native Signatur-/Startprüfung.']],[4.2,12.8],9.7)
heading('Ausdrückliche Prüfgrenzen',2)
para('Die zwölf nachfolgenden Suiten wurden wegen des Integritäts-/Fehlerstopps nicht ausgeführt: Wiederholungsregeln, Ansichten, Seitenkarten, Erwähnungen, Aufträge, Markdownbaum, Arbeitsmodi, Wiederanlauf, Speicherstatus, Sicherungszentrale, Schutzrunde und Bildkonversion. Vorhandene datierte Belege werden dafür separat bewertet; sie werden nicht in neue Testerfolge umbenannt.')
para('Die bekannte 26.420-Punkte-/40-Runden-Speicherlast, Dauer-/Hardwaretests, native Mac-/Linux-Läufe, Screenreader- und Menschenabnahmen wurden nicht neu ausgeführt. FP-01 enthält einen regulären Speicherfall mit 10.001 Punkten und 20 Undo-/Redo-Schritten; dieser lief grün. Es erfolgte keine vollständige manuelle Zeilenprüfung sämtlicher Module und keine neue Sicherheitsvollprüfung.')

matrix('3 Historischer Funktionsbestand Teil 1',data['legacy'][:8],
       'Implementierung im konsolidierten Plan dokumentiert; aktuelle Modul-/Einstiegsbelege vorhanden. „Historisch“ bezeichnet einen datierten früheren Prüfstand und keinen frischen vollständigen Regressionstest dieses Audits. Quellen E01–E05.')
matrix('4 Historischer Funktionsbestand Teil 2',data['legacy'][8:],
       'Die freigegebenen Kernumfänge bleiben von späteren Erweiterungsvorschlägen getrennt. Historische Versionsangaben bezeichnen die Einführung oder eine belegte Lieferung, nicht einen neuen nativen Abschluss von 3.49.0.')
matrix('5 Featurepakete FP 00 bis FP 07',data['fp'],
       'FP-01 bis FP-06 sowie die gewählte erste Stufe von FP-07 sind implementiert. Frühere offene Kernstufen von FP-02/03 sind durch die späteren Umsetzungen überholt. FP-01 ist aktuell wegen des roten Pflichtfalls nicht vollständig abnahmegeprüft.')
matrix('6 Arbeitsmodi und frühere Funktionsfortsetzungen',data['current'][:6],
       'D55 wählte Kompaktmodus, lesende Nebenansicht und Seiten-/Blockbedienung. Unter Windows sind vier begrenzte D55-Bedienpunkte vom 25.10.2026 dokumentiert bestätigt; das ersetzt keine globale oder plattformübergreifende Freigabe.')
matrix('7 Tagessteuerung und jüngste Dreierpakete',data['current'][6:],
       'Der gezielte Funktionsnachweis vom 28.10.2026 besitzt denselben Apphash wie die jetzt eingefrorene Quelle. Er belegt die dort genannten Bedienfälle einschließlich Wochenwahl, Fälligkeit/Aufwand und Verweisaktionen. Er ersetzt keine vollständige neue 105-Suiten-/Native-Matrix.')
matrix('8 Schutzfunktionen und technische Pakete',data['safeguards'],
       'Die begrenzten Schutzimplementierungen sind vorhanden. Insbesondere Wiederanlauf und Diagnose sind keine Reparatur der ungeklärten WA-41-Ursachen. Die neu festgestellte Prüferinkompatibilität muss in der formalen Abnahmekette beseitigt werden.')
matrix('9 Tatsächlich begonnene Restarbeiten',data['remaining'],
       'Diese Einheiten dürfen nicht pauschal als abgeschlossen geführt werden. Teilweise erfüllte technische Verträge und Piloten sind von deren appweitem Ziel getrennt. Aufwand und Reihenfolge bleiben im bestehenden Entwicklungsplan; dieser Bericht schafft keinen zweiten Backlog.')

page('10 Einzelabgleich des UI und UX Auftrags')
para('B01 bis B22 werden im Plan als 3.43.0 umgesetzt bzw. gepflegt geführt. Die Liste prüft die Vollständigkeit der Anforderungen und ihren dokumentierten Status. Sie stellt keine neue menschliche Sichtfreigabe jeder Zeile dar. Die jüngsten Designbelege umfassen 384 Knopfprüfungen in vier Größen-/Schrift-/Designkonfigurationen einschließlich Kompaktmodus und drei Scrolllagen.')
ui=next(rows for key,rows in plan.items() if key.startswith('17. UI-'))
uirows=[]
for entry in ui:
    c=entry['cells']
    if len(c)<3 or not (c[0].startswith('B') and c[0][1:].isdigit()):continue
    rest='Native/Menschenabnahme offen'
    if c[0]=='B20':rest='Richtlinien/Quellen gepflegt'
    if c[0]=='B21':rest='Dokumentpflege fortlaufend'
    if c[0]=='B22':rest='Aktuelle Lieferung/Gesamtabnahme offen'
    uirows.append([c[0],c[1],'Umgesetzt laut Plan',rest])
table(['ID','Anforderung','Implementierung','Abnahmegrenze'],uirows,[1.1,7.0,3.5,5.4],9)
para('A01–A03 sind bereits in Kapitel6 bewertet. Die Auswahl E-U2 A verlangt eine lesende Nebenansicht mit gezielten Originalaktionen; zwei gleichzeitig schreibende Editoren oder ein zweites Fenster sind kein unerledigter Rest dieser Auswahl.')

page('11 Konkrete Befunde und offene Abnahmebedingungen')
findings=[
 ('A01 Aktueller Pflichtfall für Wiederholen ist rot',
  'Der Originalfall „gueltigkeit“ scheitert nach seinem simulierten Neustart beim Anwählen eines Treeview-Elements. Die Probe zeigt davor readonly=false, danach readonly=true, Grund recovery und Sitzungsstatus active. Der Test gibt die Datensperre frei und zerstört das Fenster, bestätigt aber keinen sauberen Sitzungsabschluss. Mit zusätzlichem sauberem Abschluss ausschließlich im Diagnoseadapter besteht derselbe Fall unverändert. Die Produktionsquelle bleibt identisch. Empfehlung: Prüfharness fachgerecht anpassen, echte Abbruchfälle separat erhalten und danach Originalsuite sowie Folgematrix erneut ausführen. Bis dahin kein grüner FP-01-Gesamtnachweis. E07/E08.'),
 ('A02 P8 und mehrere Querschnittsarbeiten sind nicht vollendet',
  'Der begonnene 4.0-Aufgaben-/Readerkern ist nach D50 pausiert. WA-09, WA-32, WA-36–40 und Teile der Prüfstabilität besitzen ausdrücklich offene Restumfänge. Ein vorhandener Startseitenpilot oder Speicherkoordinator schließt nicht alle übrigen Ansichten, Caller oder Fehlerhandler ab. Die konkrete Liste steht in Kapitel9. E01/E04.'),
 ('A03 WA 41 bleibt ein eigenständiges Freigabetor',
  'WA-41A betrifft Prozess-/Interpreterabbrüche, WA-41B einen nicht aufgeklärten Save/Reload-Unterschied. Grüne AMD-Gegenproben, Schutzmarker und erfolgreiche Fachtests sind kein Ursachennachweis und keine Entwarnung für die Datenkonsistenz. Ein Datenverlust in einem echten Nutzerbestand ist nicht belegt. Keine BIOS-/CPU-Änderung oder erneute Original-Speicherlast aus diesem Audit. E03/E10.'),
 ('A04 Aktuelle Quelle ist nicht aktuell geliefert',
  'Die CI-Grundstufe meldet zehn erwartete Unterschiede zu 07_Python-Versionen. Die gelesene Python-Fassung steht auf3.47.0, beide unveränderten nativen Archive auf3.45.0. Damit sind die 3.48/3.49-Erweiterungen im Source vorhanden, aber noch nicht als vollständige aktuelle Lieferung abgenommen. E05/E06.'),
 ('A05 Dokumentation enthält überholte Einzelangaben',
  'Die manuelle Prüfliste trägt im Kopf3.49.0, nennt im ersten aktuellen Prüfabsatz aber noch „heutigen Quellkandidaten3.48.0“. Plan §18.3 führt WA-12 noch als ungeliefert, während der spätere Lieferabschluss §21.20 den umfassenderen3.45.0-Stand beschreibt. Vier Modulverweise in §18.2 existieren nicht mehr als Dateien: template_helpers.py, drawing_import.py, drawing_palette.py und drawing_animation.py. Die Funktionen sind im Integrationscode bzw. drawing.py/animation.py zu finden. Das sind Nachweis-/Pflegeprobleme, keine allein daraus bewiesenen fehlenden Features. E01/E04/E11.'),
]
for title,text in findings:
    heading(title,2);para(text)

page('12 Nicht begonnene Optionen und Freigabeplan')
para('Die folgenden Erweiterungen sind offen, jedoch keine unvollendet liegengebliebenen Implementierungen des bestätigten Kernumfangs. Ihre Planung oder Bewertung ist nicht mit einem Umsetzungsauftrag gleichzusetzen.')
table(['Vorschlag oder Tor','Einordnung'],[
    ['FP-07 JSON Canvas Stufe2','Option E-P5, nicht mit gewählter Markdownausgabe gleichsetzen'],
    ['WA-19 bis WA-25','Zusatzimporter, Wochenziele, Fundorte, Layoutbackup, Serienausnahmen/-historie und Seitenzeit: nicht pauschal ausgewählt'],
    ['WA-08/14/15, OCR/Connector/URI','Bedarfsabhängige bzw. ungewählte Erweiterungen; keine neue Umsetzung aus diesem Bericht'],
    ['WA-43/44/45/46 und WA-49/50','Release-/Wartungsaufgaben aus dem Register; Umsetzung nicht pauschal als begonnen verbuchen'],
    ['Cloud/Team/Mobile/SQL-Neubau','Keine fehlende Standardfunktion dieses Auftrags; geltende Produktgrenzen beachten']],[6.1,10.9],9.6)
heading('Vor einer Gesamtabnahme erforderlich',2)
para('1. Roten FP-01-Prüfadapter mit sauberem Sitzungsabschluss korrigieren, Originalkriterien und echte Crashfälle erhalten.\n2. Die gestoppte Folgematrix auf unveränderten Bytes durchführen und jeden Befund begründen.\n3. Begonnene Teilumfänge aus Kapitel9 abschließen oder vom Inhaber ausdrücklich aus der jeweiligen Teilabnahme ausnehmen.\n4. WA-41A/B und aktuelle breite/native Prüfgrenzen gesondert klären.\n5. Python/Showcase und eigenständige Pakete auf den freigegebenen Quellstand bringen; Hashes, native Starts und Signaturnachweise prüfen.\n6. Mac/Windows/Linux, DPI/Mehrmonitor, Reader/Tastatur und die offenen Menschenpunkte aus der bestehenden Prüfliste abnehmen. Öffentliche Lizenz-/Signierungs-/Lieferkettentore bleiben getrennt.')
heading('Dokumentation der Entscheidung',2)
para('[ ] Technische Teilabnahme nur der ausdrücklich benannten Implementierungen\n[ ] Gesamtabnahme zurückgestellt\n[ ] Erneute Prüfung nach Erfüllung der offenen Bedingungen')
para('Abgenommener Teilumfang und Ausnahmen:\n______________________________________________________________________\n______________________________________________________________________')
para('Abnehmende Person _____________________   Datum _____________________\nBestätigung / Unterschrift _____________________________________________')

page('13 Quellen und unverwechselbarer Prüfstand')
sources=[
 ['E01','00_Arbeitsvorbereitung/Glide_Entwicklungsplan.md','§§15.7,17,18,19,21.1–21.28; Kennungen/Scope/Status'],
 ['E02','01_Repository/Glide/docs/20_FUNKTIONEN.md','Verhalten der bestehenden und neuen Funktionen'],
 ['E03','01_Repository/Glide/docs/07_QA_BERICHT.md','Datierte Prüf-/Liefergrenzen, WA-41'],
 ['E04','00_Arbeitsvorbereitung/Glide_Manuelle_Pruefung.md','Offene Menschenabnahme und vier Windows-Teilbestätigungen'],
 ['E05','lieferstand.json','Neu gelesene Versionen/Metadaten und Archivhashvergleich'],
 ['E06','lauf_CI-Grundstufe.json und testzaehlung.json','Neuer vollständiger Grundstufenlauf auf eingefrorener Quelle'],
 ['E07','lauf_wiederholen.json und testergebnis.json','Originaler roter Pflichtfall; gestoppte Folge'],
 ['E08','diagnose_wiederholen.json','Reproduktion und isolierte Lifecyclevariation ohne Quell-/Assertionsänderung'],
 ['E09','../funktionspakete_2026-10-28/ergebnis.json','Gezielte D57-Bedienbelege; Apphash aktuell identisch'],
 ['E10','../robustheit_2026-10-28/README.md','Historische Schutz-/Restore-/Diagnosebelege, eigene Grenzen'],
 ['E11','codeinventar.json, planregister.json, paketmatrix.json','Modul-/Einstiegsbelege, Statusauszug und bewertete Abnahmeeinheiten'],
 ['E12','eingefroren.json und testergebnis.json','Quellhashes, Laufzeit, Unverändertheit und lokale Rohlogorte'],
]
table(['Beleg','Quelle','Verwendung'],sources,[1.2,8.2,7.6],8.7)
para('E05–E12 liegen, soweit kein anderer relativer Pfad genannt ist, neben diesem Bericht. Die JSON-Dateien enthalten die detaillierten maschinellen Belege. Rohlogs und künstliche Prüfdaten bleiben in der im Manifest benannten lokalen Ablage. Kein Bericht führt eine menschliche Unterschrift oder einen aktuellen GitHub-/Merge-/Releaseentscheid herbei.')
para('Maßgeblicher App-SHA-256:',True)
p=para(freeze['app_sha256']);p.runs[0].font.name='Consolas';p.runs[0].font.size=Pt(8.5)
para('Letzter lokaler Commit: '+freeze['head']+' (3.47.0-Basis). Der aktuelle3.49.0-Arbeitsstand ist uncommittet; ein Checkout allein dieses Commits ist nicht der geprüfte Kandidat.')
para(f"Strukturinventar: {code['app_lines']:,} Zeilen in app.pyw und {len(code['listapp_methods']):,} direkte ListApp-Methoden; 89 .py-Dateien direkt im Modulordner einschließlich Startmodul und Zeichenprototyp, app.pyw zusätzlich. Diese Größenwerte belegen den Wartungsumfang, keinen konkreten Laufzeitfehler.".replace(',', '.'))
para('Anwendungscode, Nutzerdaten und Lieferartefakte wurden durch dieses Audit nicht verändert. Keine neue Appversion, keine Hardwareänderung, kein Commit/Push, kein Paketbau und keine Releasefreigabe. Die Prüfung ist ein technischer Abnahmebericht mit offen ausgewiesenen Grenzen, kein vollständiger Zeilen-/Sicherheitsaudit der gesamten Codebasis.')
doc.save(DOCX)
print(DOCX)
