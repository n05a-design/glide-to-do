from __future__ import annotations

import json
import re
import sys
import unicodedata
import zipfile
from collections import Counter
from difflib import SequenceMatcher
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from pypdf import PdfReader


DOCX = Path(sys.argv[1])
PDF = Path(sys.argv[2])


def norm(value: str) -> str:
    value = unicodedata.normalize("NFKC", value).casefold()
    value = value.replace("\u00ad", "")
    value = re.sub(r"[^\w]+", " ", value, flags=re.UNICODE)
    return re.sub(r"\s+", " ", value).strip()


def docx_text(doc: Document) -> str:
    chunks: list[str] = []
    for section in doc.sections:
        for area in (section.header, section.footer):
            chunks.extend(p.text for p in area.paragraphs if p.text.strip())
            for table in area.tables:
                for row in table.rows:
                    for cell in row.cells:
                        chunks.extend(p.text for p in cell.paragraphs if p.text.strip())
    for p in doc.paragraphs:
        if p.text.strip():
            chunks.append(p.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                chunks.extend(p.text for p in cell.paragraphs if p.text.strip())
    return "\n".join(chunks)


doc = Document(DOCX)
reader = PdfReader(str(PDF))
pdf_pages = [(page.extract_text() or "") for page in reader.pages]
doc_text = docx_text(doc)
pdf_text = "\n".join(pdf_pages)
doc_tokens = norm(doc_text).split()
pdf_tokens = norm(pdf_text).split()
doc_counts = Counter(doc_tokens)
pdf_counts = Counter(pdf_tokens)
shared = sum((doc_counts & pdf_counts).values())
missing_from_pdf = list((doc_counts - pdf_counts).elements())
extra_in_pdf = list((pdf_counts - doc_counts).elements())
precision = shared / max(len(pdf_tokens), 1)
recall = shared / max(len(doc_tokens), 1)
ordered_ratio = SequenceMatcher(None, norm(doc_text), norm(pdf_text), autojunk=False).ratio()

ns = {
    "ep": "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties",
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
}
with zipfile.ZipFile(DOCX) as zf:
    names = set(zf.namelist())
    app_root = ET.fromstring(zf.read("docProps/app.xml"))
    app = {}
    for key in ("Application", "AppVersion", "Pages", "Words", "Characters", "Paragraphs", "Lines"):
        elem = app_root.find(f"ep:{key}", ns)
        app[key] = elem.text if elem is not None else None
    document_xml = zf.read("word/document.xml")
    root = ET.fromstring(document_xml)
    tracked = {
        "insertions": len(root.findall(".//w:ins", ns)),
        "deletions": len(root.findall(".//w:del", ns)),
        "moves_from": len(root.findall(".//w:moveFrom", ns)),
        "moves_to": len(root.findall(".//w:moveTo", ns)),
    }
    comments = [n for n in names if re.fullmatch(r"word/comments\d*\.xml", n)]

page_summaries = []
for i, text in enumerate(pdf_pages, 1):
    collapsed = re.sub(r"\s+", " ", text).strip()
    page_summaries.append({
        "page": i,
        "characters": len(text),
        "start": collapsed[:220],
        "end": collapsed[-220:],
    })

props = doc.core_properties
result = {
    "docx_path": str(DOCX),
    "pdf_path": str(PDF),
    "docx_core_properties": {
        "title": props.title,
        "subject": props.subject,
        "author": props.author,
        "last_modified_by": props.last_modified_by,
        "created": props.created.isoformat() if props.created else None,
        "modified": props.modified.isoformat() if props.modified else None,
        "revision": props.revision,
        "keywords": props.keywords,
        "comments": props.comments,
    },
    "docx_extended_properties": app,
    "pdf_pages": len(reader.pages),
    "pdf_metadata": {str(k): str(v) for k, v in (reader.metadata or {}).items()},
    "docx_text_characters": len(doc_text),
    "pdf_text_characters": len(pdf_text),
    "docx_tokens": len(doc_tokens),
    "pdf_tokens": len(pdf_tokens),
    "shared_token_occurrences": shared,
    "missing_docx_tokens_in_pdf": missing_from_pdf[:100],
    "extra_pdf_tokens_not_in_docx_sample": extra_in_pdf[:100],
    "token_precision_pdf_vs_docx": round(precision, 6),
    "token_recall_docx_vs_pdf": round(recall, 6),
    "ordered_character_similarity": round(ordered_ratio, 6),
    "tracked_changes": tracked,
    "comment_parts": comments,
    "page_summaries": page_summaries,
}
print(json.dumps(result, ensure_ascii=False, indent=2))
