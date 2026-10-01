from __future__ import annotations

import argparse
import re
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


TOC_BOOKMARK_PREFIX = "GlideToc"
DATA_TABLE_INDICES = (4, 6, 7, 9, 10, 11, 12, 14)


def clear_paragraph(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_text(paragraph, text: str, *, bold_prefix: str | None = None) -> None:
    clear_paragraph(paragraph)
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        first.bold = True
        paragraph.add_run(text[len(bold_prefix) :])
    else:
        paragraph.add_run(text)


def find_paragraph(document, text: str):
    for paragraph in document.paragraphs:
        if paragraph.text == text:
            return paragraph
    raise KeyError(f"Absatz nicht gefunden: {text}")


def find_paragraph_start(document, prefix: str):
    for paragraph in document.paragraphs:
        if paragraph.text.startswith(prefix):
            return paragraph
    raise KeyError(f"Absatz nicht gefunden: {prefix}")


def insert_after(document, anchor, text: str, style: str = "Normal"):
    paragraph = document.add_paragraph(text, style=style)
    anchor._p.addnext(paragraph._p)
    return paragraph


def set_cell_text(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    set_text(paragraph, text)
    for extra in list(cell.paragraphs[1:]):
        extra._element.getparent().remove(extra._element)


def set_callout(table, heading: str, body: str) -> None:
    cell = table.cell(0, 0)
    set_text(cell.paragraphs[0], heading)
    cell.paragraphs[0].runs[0].bold = True
    set_text(cell.paragraphs[1], body)


def iter_story_paragraphs(document):
    for paragraph in document.paragraphs:
        yield paragraph
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in document.sections:
        for story in (
            section.header,
            section.footer,
            section.first_page_header,
            section.first_page_footer,
            section.even_page_header,
            section.even_page_footer,
        ):
            yield from story.paragraphs
            for table in story.tables:
                for row in table.rows:
                    for cell in row.cells:
                        yield from cell.paragraphs


def replace_in_stories(document, old: str, new: str) -> None:
    for paragraph in iter_story_paragraphs(document):
        for run in paragraph.runs:
            if old in run.text:
                run.text = run.text.replace(old, new)


def ensure_rfonts(rpr, font_name: str) -> None:
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attribute}"), font_name)


def apply_cross_renderer_fonts(document) -> None:
    for style in document.styles:
        if style.type not in (WD_STYLE_TYPE.PARAGRAPH, WD_STYLE_TYPE.CHARACTER):
            continue
        font_name = "Consolas" if style.name == "CodeBlock" else "Arial"
        style.font.name = font_name
        rpr = style.element.get_or_add_rPr()
        ensure_rfonts(rpr, font_name)

    if "Normal" in document.styles:
        document.styles["Normal"].font.name = "Arial"
        document.styles["Normal"].font.size = Pt(9.5)
    if "CodeBlock" in document.styles:
        document.styles["CodeBlock"].font.name = "Consolas"
        document.styles["CodeBlock"].font.size = Pt(8.0)

    for paragraph in iter_story_paragraphs(document):
        style_name = paragraph.style.name if paragraph.style is not None else ""
        font_name = "Consolas" if style_name == "CodeBlock" else "Arial"
        for run in paragraph.runs:
            run.font.name = font_name
            ensure_rfonts(run._r.get_or_add_rPr(), font_name)


def add_or_replace_child(parent, tag: str, **attributes):
    for existing in list(parent.findall(qn(tag))):
        parent.remove(existing)
    element = OxmlElement(tag)
    for key, value in attributes.items():
        element.set(qn(key), value)
    parent.append(element)
    return element


def format_data_tables(document) -> None:
    for table_index in DATA_TABLE_INDICES:
        table = document.tables[table_index]
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        for row_index, row in enumerate(table.rows):
            tr_pr = row._tr.get_or_add_trPr()
            add_or_replace_child(tr_pr, "w:cantSplit")
            if row_index == 0:
                add_or_replace_child(tr_pr, "w:tblHeader", **{"w:val": "true"})
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.keep_with_next = True


def add_bookmark(paragraph, bookmark_id: int, name: str) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def add_hyperlink_run(hyperlink, text: str, *, tab: bool = False) -> None:
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{attribute}"), "Arial")
    r_pr.append(r_fonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "17")
    r_pr.append(size)
    size_cs = OxmlElement("w:szCs")
    size_cs.set(qn("w:val"), "17")
    r_pr.append(size_cs)
    run.append(r_pr)
    if tab:
        run.append(OxmlElement("w:tab"))
    else:
        node = OxmlElement("w:t")
        node.text = text
        run.append(node)
    hyperlink.append(run)


def add_internal_toc_link(paragraph, text: str, bookmark: str, page: str = "00") -> None:
    clear_paragraph(paragraph)
    paragraph.paragraph_format.space_after = Pt(1.5)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        Cm(7.55),
        WD_TAB_ALIGNMENT.RIGHT,
        WD_TAB_LEADER.DOTS,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), bookmark)
    hyperlink.set(qn("w:history"), "1")
    add_hyperlink_run(hyperlink, text)
    add_hyperlink_run(hyperlink, "", tab=True)
    add_hyperlink_run(hyperlink, page)
    paragraph._p.append(hyperlink)


def rebuild_toc(document) -> None:
    headings = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.style is not None
        and paragraph.style.name == "Heading 1"
        and re.match(r"^\d+\.\s", paragraph.text)
    ]
    if len(headings) != 23:
        raise ValueError(f"Erwartet wurden 23 Hauptkapitel, gefunden: {len(headings)}")

    entries = []
    for index, heading in enumerate(headings, start=1):
        bookmark = f"{TOC_BOOKMARK_PREFIX}{index:02d}"
        add_bookmark(heading, 1000 + index, bookmark)
        entries.append((heading.text, bookmark))

    toc_table = document.tables[1]
    columns = (entries[:12], entries[12:])
    for cell, cell_entries in zip(toc_table.rows[0].cells, columns):
        first = cell.paragraphs[0]
        for extra in list(cell.paragraphs[1:]):
            extra._element.getparent().remove(extra._element)
        for entry_index, (label, bookmark) in enumerate(cell_entries):
            paragraph = first if entry_index == 0 else cell.add_paragraph()
            add_internal_toc_link(paragraph, label, bookmark)


def format_pagination(document) -> None:
    for paragraph in document.paragraphs:
        style_name = paragraph.style.name if paragraph.style is not None else ""
        if style_name.startswith("Heading"):
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.keep_together = True
        if style_name == "CodeBlock":
            line_count = paragraph.text.count("\n") + 1
            if line_count <= 22:
                paragraph.paragraph_format.keep_together = True

    agents_block = find_paragraph_start(document, "# Glide – Codex Instructions")
    agents_block.paragraph_format.page_break_before = True
    agents_block.paragraph_format.keep_together = True

    phase_heading = find_paragraph(document, "19. Verbindlicher Phasenplan")
    phase_heading.paragraph_format.page_break_before = True

    source_heading = find_paragraph(document, "23. Quellen und Prüfstand")
    source_seen = False
    for paragraph in document.paragraphs:
        if paragraph is source_heading:
            source_seen = True
            continue
        if source_seen and paragraph.style is not None and paragraph.style.name == "List Bullet":
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            for run in paragraph.runs:
                run.font.size = Pt(9)


def format_footer(document) -> None:
    # The intermediate document keeps the source footer content. The release
    # finalizer replaces this table with a renderer-stable, tab-aligned line
    # after the final page numbers have been determined.
    for section in document.sections:
        footer = section.footer
        if footer.tables:
            table = footer.tables[0]
            table.autofit = False
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for column in table.columns:
                column.width = Cm(5.48)
            for cell in table.rows[0].cells:
                cell.width = Cm(5.48)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = "Arial"
                        run.font.size = Pt(8)


def update_document(source: Path, output: Path) -> None:
    document = Document(source)

    set_text(find_paragraph_start(document, "Dokumentversion 2.5.1"), "Dokumentversion 2.5.2 · Stand 01.09.2026")
    set_text(
        find_paragraph_start(document, "Arbeitsstatus:"),
        "Arbeitsstatus: Source 2.5.2 implementiert, migriert und getestet · Packaging/Release weiterhin Pre-Release",
    )
    replace_in_stories(document, "Stand 31.08.2026", "Stand 01.09.2026")

    set_callout(
        document.tables[0],
        "Zweck und Abgrenzung dieses Dokuments",
        "Diese Revision dokumentiert den tatsächlich erreichten Source-, Test- und Strukturstand von Glide 2.5.2 sowie die weiterhin offenen Build-, Signing- und Store-Arbeiten. Historische Screenshots und Ticketvorlagen sind Referenzmaterial, keine eigenständigen Ausführungsanweisungen.",
    )
    set_callout(
        document.tables[2],
        "Kernaussage 2.5.2",
        "Der Eingang ist sichtbar von den normalen Listen getrennt; Aufgaben besitzen ein direktes Kontextmenü für Bearbeitung, Beschreibung, Anhänge, Wichtigkeit, Fälligkeit und Farbe. Der gespeicherte Dark Mode färbt die native Windows-Titelleiste bereits beim Start. Datenformat 5 und Format-4-Komplettbackups sind automatisiert geprüft.",
    )

    set_text(find_paragraph(document, "Inhaltsverzeichnis"), "Inhaltsübersicht")
    set_text(
        find_paragraph_start(document, "Die Kapitel sind in Prozessreihenfolge"),
        "Die Kapitel sind in Prozessreihenfolge angeordnet. Die Seitenzahlen und internen Sprungziele werden aus dem finalen Render dieser Revision erzeugt.",
    )

    set_text(
        find_paragraph_start(document, "Gesamtbewertung."),
        "Gesamtbewertung. Glide 2.5.2 liegt als lauffähiger, rückwärtskompatibler und auf zwei lokalen Python/Tk-Laufzeiten getesteter Source-Stand vor. Zusätzlich zu den 2.5.1-Funktionen sind der getrennte Eingang, das Aufgaben-Kontextmenü, persistente Aufgabenfarben und der zuverlässige Windows-Dark-Start umgesetzt. Der vorherige Lauf war nicht im Code abgebrochen, sondern unmittelbar nach Word-Erzeugung und 21-Seiten-Render vor der visuellen Dokumentprüfung.",
        bold_prefix="Gesamtbewertung.",
    )
    set_text(
        find_paragraph_start(document, "Der nächste Qualitätssprung liegt nun"),
        "Der nächste Qualitätssprung liegt weiterhin außerhalb der reinen Feature-Implementierung: vollständige manuelle Plattform-QA, verbindliche Produktidentitäten, reproduzierbares Packaging, Signing/Notarisierung und Clean-Machine-Tests. Die bestehende Struktur trennt dafür Source, Tests, Fixtures, technische Dokumentation, Grafik-Master, Store-Material und finale Release-Artefakte.",
    )

    set_text(
        find_paragraph_start(document, "Abbildung 1:"),
        "Abbildung 1: Historischer Zwischenstand während des vorherigen Laufs (31.08.2026).",
    )
    historical = find_paragraph_start(document, "Der Screenshot zeigt den Zustand")
    set_text(
        historical,
        "Der Screenshot dokumentiert einen frühen Zwischenstand des vorherigen Laufs, nicht dessen tatsächlichen Abbruchpunkt. Zu diesem Zeitpunkt lagen Versionsdateien unter 07_Python-Versionen, der Test noch im Workspace-Root und ein generierter __pycache__ daneben. Danach wurden Source 2.5.1, Tests, Struktur und Word-Dokument bereits fertiggestellt.",
    )

    set_text(find_paragraph(document, "2.4 Umgesetzter Arbeitsstand 2.5.1"), "2.4 Umgesetzter Arbeitsstand 2.5.2")
    set_text(
        find_paragraph_start(document, "Die neue Struktur wurde nicht-destruktiv angelegt."),
        "Die Struktur bleibt nicht-destruktiv. Der kanonische Source wurde auf 2.5.2 fortgeführt; 2.5.1 bleibt als unveränderte historische Einzeldatei erhalten, und die neue 2.5.2-Einzeldatei ist bytegleich zum kanonischen Source.",
    )
    structure = find_paragraph_start(document, "Glide_Workspace/")
    set_text(
        structure,
        "Glide_Workspace/\n│\n├─ 00_Arbeitsvorbereitung/{Entscheidungen, Research, Checklisten, Notizen}/\n├─ 01_Repository/Glide/\n│  ├─ AGENTS.md · README.md · CHANGELOG.md · VERSION\n│  ├─ src/glide/{app.pyw, resources/}\n│  ├─ tests/{unit, integration, smoke, packaging}/\n│  ├─ tests/fixtures/{current_v5, current_v4, legacy_v2}/\n│  ├─ assets/ · packaging/ · scripts/ · requirements/\n│  └─ docs/\n│     ├─ 00_INDEX.md · 01_PRODUCT_CONSTRAINTS.md · 02_ARCHITECTURE.md\n│     ├─ 05_QA_TESTPLAN.md · 06_DATA_BACKUP_MIGRATION.md\n│     └─ 10_RELEASE_CHECKLIST.md · decisions/ · exec-plans/\n├─ 05_Probelisten_Testdaten/          (Legacy-Quelle)\n├─ 07_Python-Versionen/               (Versionshistorie bis 2.5.2)\n├─ 10_Dokumentation/                  (Word/PDF)\n├─ 20_Grafik_Master/\n├─ 30_Release_Exports/{2.5.1, 2.5.2}/ (noch ohne Build)\n├─ 40_Store_Material/\n├─ 90_Testdaten_Extern/\n└─ 100_Archiv/",
    )
    root_test = find_paragraph_start(document, "Der frühere Root-Test wurde")
    set_text(
        root_test,
        "Der kanonische Integrationstest prüft nun zusätzlich getrennten Eingang, Aufgaben-Kontextmenü und -farben, Format-4-Backups, die eingecheckten v2/v4/v5-Fixtures sowie den realen Windows-DWM-Wert. Es wurde weiterhin bewusst kein Git-Repository initialisiert und kein Packaging- oder Signingwert erfunden.",
    )
    anchor = insert_after(document, root_test, "2.5 Rekonstruierter Abbruchpunkt und nachgeholte Dokument-QA", "Heading 2")
    anchor = insert_after(
        document,
        anchor,
        "Die eingebetteten Zeitangaben belegen die Reihenfolge: Source 2.5.1 und Tests waren fertig; die DOCX wurde am 31.08.2026 um 18:21:39 MESZ geändert und das zugehörige 21-seitige PDF 27 Sekunden später erzeugt. Danach fehlt ein abgeschlossener QA-Bericht. Der Lauf endete somit unmittelbar vor der angekündigten visuellen Kontrolle und finalen Übergabe.",
    )
    insert_after(
        document,
        anchor,
        "Die 21 Seiten wurden in dieser Revision vollständig nachgeprüft. Es gab keinen Inhaltsverlust und keine beschädigten Grafiken, aber fehlende Tabellenkopf-Wiederholungen, teilbare Zeilen, einen verwaisten AGENTS-Codeblock, ein statisches Inhaltsverzeichnis und rendererabhängige Aptos-Ersatzschriften. Die 2.5.2-Fassung korrigiert diese Punkte mit wiederholten Tabellenköpfen, unteilbaren Zeilen, kontrollierten Seitenumbrüchen, internen Sprungzielen und expliziten Arial-/Consolas-Schriften.",
    )

    set_text(
        find_paragraph_start(document, "Die empfohlene Zweiteilung ist für 2.5.1"),
        "Die Zweiteilung ist für 2.5.2 fortgeführt: ein persönlicher Workspace für Arbeitsmaterialien und ein technisch kontrollierter Repository-Bereich für Source, Tests, Packaging-Vorbereitung und code-nahe Dokumentation.",
    )
    set_callout(
        document.tables[8],
        "Wichtig für den nächsten Schritt",
        "Der funktionsfähige Monolith ist als 2.5.2 gesichert und geprüft; kanonischer Source und historische 2.5.2-Einzeldatei sind bytegleich. Eine inhaltliche Modulzerlegung bleibt ein separates Refactoring und darf erst mit zusätzlichen Regressionstests erfolgen.",
    )

    set_text(find_paragraph_start(document, "C01 – Repository Audit"), "C01 – Repository Audit (für 2.5.2 abgeschlossen)")

    identity = find_paragraph_start(document, "Datei:\n01_Repository/Glide/docs/decisions/PRODUCT_IDENTITY.md")
    set_text(
        identity,
        "Datei:\n01_Repository/Glide/docs/decisions/PRODUCT_IDENTITY.md\n\nProduktname:\nGlide\n\nLangname:\nGlide – Aufgaben und Listen\n\nAktuelle App-Version:\n2.5.2\n\nDatenformat-Version:\n5\n\nRelease-Kanal:\ninternal / pre-release\n\nPublisher, Copyright, Support-E-Mail, Website, Datenschutz-URL und Lizenzmodell:\noffen – nicht durch Annahmen ersetzen\n\nWindows:\nZiel Windows 10/11; Architektur, AppUserModelID und Inno-Setup-AppId final festlegen\n\nmacOS:\nBundle Identifier, Mindestversion und Architektur final festlegen",
    )

    set_text(
        find_paragraph_start(document, "Die QA-Grundlage ist jetzt operationalisiert:"),
        "Die QA-Grundlage ist erweitert: Der Integrationstest startet Glide mit isoliertem temporärem Datenordner, verarbeitet echte Tk-Ereignisse und prüft Kernfunktionen, UI-Metadaten, Migration, Backup/Restore sowie mehrere Fehlerfälle. Unter Windows wird der DWM-Dark-Wert beim ersten Mapping und beim Theme-Wechsel ausgelesen. Unit-, Packaging- und reale Clean-Machine-Tests bleiben als nächste Ebenen offen.",
    )
    test_tree = find_paragraph_start(document, "tests/\n├─ unit/")
    set_text(
        test_tree,
        "tests/\n├─ unit/                         (vorbereitet)\n├─ integration/test_glide.py    (grün)\n├─ smoke/                        (vorbereitet)\n├─ packaging/                    (vorbereitet)\n└─ fixtures/\n   ├─ current_v5/reference_v5.json\n   ├─ current_v4/reference_v4.json\n   └─ legacy_v2/probelisten_5_listen_v2.json",
    )
    set_text(
        find_paragraph_start(document, "Nachweis am 31.08.2026:"),
        "Nachweis am 01.09.2026: „Glide v2.5.2 Kern-, Backup-, UI- und Migrationstests: OK“ auf Python 3.13.15 sowie der gebündelten Python-3.12/Tk-8.6-Laufzeit. Die nachfolgende Matrix beschreibt Sollprüfungen; „erforderlich“ bedeutet nicht automatisch „bestanden“.",
    )

    matrix = document.tables[12]
    set_cell_text(matrix.cell(0, 1), "Windows – Soll")
    set_cell_text(matrix.cell(0, 2), "macOS – Soll")
    for row in matrix.rows[1:]:
        for cell in row.cells[1:]:
            value = cell.text.strip()
            if value == "Ja":
                set_cell_text(cell, "erforderlich")
            elif value == "–":
                set_cell_text(cell, "nicht anwendbar")
            elif value == "Retina/Skalierung":
                set_cell_text(cell, "erforderlich (Retina)")

    set_text(
        find_paragraph_start(document, "Die Referenzdaten sind nun getrennt:"),
        "Die Referenzdaten sind dreistufig: ein aktuelles Format-5-Fixture, das frühere Format-4-Fixture für die neue Migration sowie ein größerer Format-2-Bestand. Alle drei werden vom Integrationstest tatsächlich geladen und normalisiert. Manuelle TXT-Importbeispiele bleiben außerhalb der automatisierten Tests.",
    )
    fixture_tree = find_paragraph_start(document, "01_Repository/Glide/tests/fixtures/")
    set_text(
        fixture_tree,
        "01_Repository/Glide/tests/fixtures/\n├─ current_v5/reference_v5.json\n├─ current_v4/reference_v4.json\n└─ legacy_v2/probelisten_5_listen_v2.json\n\n90_Testdaten_Extern/Legacy_Probelisten/\n└─ manuelle TXT-Importbeispiele",
    )
    set_text(
        find_paragraph_start(document, "Das aktuelle Fixture enthält Eingang"),
        "Das Format-5-Fixture enthält getrennten Eingang, Ordner, Seitennotiz, Beschreibung, Fälligkeit, Wichtigkeit, Aufgabenfarbe und Unterpunkt. Format 4 prüft die fehlende Farbe als rückwärtskompatible Migration; der größere Altbestand bleibt ausdrücklich Legacy-v2-Fixture.",
    )

    set_text(
        find_paragraph_start(document, "Finale Veröffentlichungen werden ausschließlich"),
        "Finale Veröffentlichungen werden ausschließlich außerhalb des Repository-Bereichs unter 30_Release_Exports/<VERSION>/ abgelegt. Für 2.5.2 ist die Struktur vorbereitet, aber noch kein EXE-, Installer-, App- oder DMG-Artefakt als freigegeben vorhanden. Das folgende Schema bleibt eine Vorlage.",
    )

    set_text(
        find_paragraph_start(document, "Abgeschlossen: Der lauffähige Source-Stand 2.5.1"),
        "Abgeschlossen: Der lauffähige Source-Stand 2.5.2 ist unter 07_Python-Versionen erhalten und bytegleich im kanonischen Repository-Pfad vorhanden. 2.5.1 bleibt unverändert als vorherige stabile Revision bestehen.",
    )
    set_text(find_paragraph_start(document, "Glide_PreRelease_2.5.1"), "Glide_PreRelease_2.5.2_2026-09-01")
    doc_tree = find_paragraph_start(document, "AGENTS.md                              vorhanden")
    set_text(
        doc_tree,
        "AGENTS.md                              vorhanden\ndocs/\n├─ 00_INDEX.md                         vorhanden\n├─ 01_PRODUCT_CONSTRAINTS.md           vorhanden\n├─ 02_ARCHITECTURE.md                  auf 2.5.2 aktualisiert\n├─ 05_QA_TESTPLAN.md                   auf 2.5.2 aktualisiert\n├─ 06_DATA_BACKUP_MIGRATION.md         auf Format 5 aktualisiert\n├─ 10_RELEASE_CHECKLIST.md             auf 2.5.2 aktualisiert\n├─ decisions/PRODUCT_IDENTITY.md       Version 2.5.2 / Format 5\n├─ exec-plans/2.5.1-stabilisierung.md  historische Revision\n├─ exec-plans/2.5.2-qol-stabilisierung.md vorhanden\n└─ Build/Branding/Signing/Store-Details folgen mit den jeweiligen Arbeitspaketen",
    )

    phase = document.tables[14]
    phase_updates = {
        1: "teilweise abgeschlossen: Source 2.5.2, Fixtures, Dokumentation und QA-Grundlage vorhanden; Identität und Branding offen",
        2: "für 2.5.2 abgeschlossen: Source, Abhängigkeiten, Pfade, Daten, Risiken und tatsächlicher Abbruchpunkt geprüft",
        3: "teilweise abgeschlossen: Struktur, AGENTS, Doku und Integrationstest vorhanden; Git-Initialisierung und Modulzerlegung offen",
        7: "teilweise: zwei lokale Python/Tk-Läufe sowie Windows-UI/DWM-Smoke grün; vollständige Plattform-, Packaging- und Clean-Machine-QA offen",
    }
    for row_index, value in phase_updates.items():
        set_cell_text(phase.cell(row_index, 2), value)

    set_text(
        find_paragraph_start(document, "Der zuvor unterbrochene Entwicklungsstand ist als Glide 2.5.1"),
        "Der unterbrochene Dokument-QA-Schritt ist nachgeholt und der Source als Glide 2.5.2 fortgeführt. Getrennter Eingang, Aufgaben-Kontextmenü, persistente Farben und Dark-Start sind umgesetzt und geprüft; verbleibende Risiken liegen vor allem in vollständiger manueller Plattform-QA, Produktidentität, Branding, Packaging, Signing, Notarisierung und Store-Prozessen.",
    )
    set_text(find_paragraph(document, "Empfohlene Reihenfolge nach Abschluss von 2.5.1:"), "Empfohlene Reihenfolge nach Abschluss von 2.5.2:")
    set_text(
        find_paragraph_start(document, "1. Den neuen 2.5.1-Stand"),
        "1. Den neuen 2.5.2-Stand mit einer Kopie realer Arbeitsdaten manuell prüfen; zuvor ein 2.5.1-Komplettbackup erzeugen.",
    )
    set_text(
        find_paragraph_start(document, "2. Windows 10/11:"),
        "2. Die vollständige Windows-10-/Windows-11-Matrix für NumLock, CapsLock, Dark-Start, Menüs, Dialoge, Kontextaktionen und Drag & Drop dokumentieren.",
    )
    set_callout(
        document.tables[16],
        "Nächster sinnvoller Arbeitsauftrag",
        "Nach fachlicher Prüfung von 2.5.2 sollte C03 als klar begrenztes Packaging-/Release-Härtungs-Ticket folgen: zuerst Produktidentitäten festlegen, dann einen reproduzierbaren unsigned PyInstaller-OneDir-Build und passende Smoke Tests erstellen. C04–C08 bleiben separate Aufträge.",
    )

    set_text(
        find_paragraph_start(document, "Glide-Source v2.4.1"),
        "Glide-Source v2.4.1, angefangener Zwischenstand v2.5.0, stabilisierter Source v2.5.1 und fortgeführter Source v2.5.2.",
    )
    set_text(
        find_paragraph_start(document, "Lokaler Workspace, neue Struktur"),
        "Lokaler Workspace, Struktur unter 01_Repository/Glide, beide grüne Python/Tk-Testläufe und sichtbarer Windows-UI/DWM-Smoke für v2.5.2.",
    )
    set_text(
        find_paragraph_start(document, "Prüfstand: 31.08.2026."),
        "Prüfstand: 01.09.2026. Source, Struktur, Referenzdaten und automatisierter Integrationstest entsprechen App-Version 2.5.2 und Datenformat 5; Format-4-Komplettbackups bleiben unterstützt. Store-, Signierungs- und Plattformanforderungen wurden in dieser Feature-Revision nicht neu recherchiert und müssen unmittelbar vor Einreichung oder Zertifikatsbeschaffung anhand offizieller Quellen erneut geprüft werden.",
    )

    format_data_tables(document)
    rebuild_toc(document)
    format_pagination(document)
    apply_cross_renderer_fonts(document)
    format_footer(document)

    properties = document.core_properties
    properties.title = "Glide 2.5.2 – Arbeitsvorbereitung für Codex, Build, Release und Dokumentation"
    properties.subject = "Source-Stand 2.5.2, Dokument-QA, Datenformat 5 und verbleibender Release-Plan"
    properties.keywords = "Glide, 2.5.2, Tkinter, QA, Dark Mode, Aufgabenfarbe, Backup, Migration, Release"
    properties.comments = "Fortgeführte Revision mit rekonstruierter Abbruchstelle, nachgeholter Seitenprüfung und 2.5.2-Funktionsstand."
    properties.last_modified_by = "OpenAI Codex"
    properties.revision = 3
    properties.modified = datetime.now(timezone.utc).replace(tzinfo=None)

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    arguments = parser.parse_args()
    update_document(arguments.source, arguments.output)


if __name__ == "__main__":
    main()
